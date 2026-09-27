import sys
from datetime import date
from pathlib import Path

import openpyxl
import pytest

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))


def _build_workbook(tmp_path):
    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    pa = wb.create_sheet("PA")
    pa.append(["Product Name", "Category", "MRP", "Cost", "SP"])
    pa.append(["Tata Salt", "Grocery", 28, 22, 25])
    pa.append(["Tata Salt", "Grocery", 30, 24, 27])  # later dup -> should win
    pa.append(["Parle G", "Biscuits", 10, 7, 9])
    pa.append([None, None, None, None, None])  # blank -> skipped
    pa.append(["Bad Item", "Grocery", "abc", 5, 5])  # non-numeric -> error
    pa.append(["Negative Item", "Grocery", -5, 5, 5])  # negative -> error

    sales = wb.create_sheet("Sales Register")
    sales.append(["header row 1 placeholder"])
    sales.append(["Date", "CUSTOMER NO.", "Item", "Qty", "sold(INR)", "Amount", "Payment Mode",
                   "IF CREDIT", "DAYS TOTAL", "SHIFT", "PROFIT", "AVERAGE"])
    sales.append([date(2026, 7, 1), "Grocery", "Tata Salt", 2, 25, 50, "Cash", None, None, "MORNING", None, None])
    sales.append([date(2026, 7, 1), "Grocery", "Parle G", 1, 9, 9, "Cash", "credit", None, "MORNING", None, None])
    sales.append([date(2026, 7, 2), None, "Tata Salt", 1, 27, 27, "Credit", "Ramesh", None, "EVENING", None, None])
    sales.append([date(2026, 7, 2), None, None, "CASH TOTAL", None, 500, None, None, None, None, None, None])
    sales.append([date(2026, 7, 3), None, "Incomplete Item", None, None, 0, "Cash", None, None, None, None, None])
    sales.append([None, None, None, None, None, None, None, None, None, None, None, None])

    stock = wb.create_sheet("Stock Register")
    stock.append(["Item", "Something", "Opening"])
    stock.append(["Tata Salt", None, 50])
    stock.append(["Parle G", None, 30])
    stock.append(["New Widget", None, 10])  # missing from PA -> auto-created

    purch = wb.create_sheet("Purchase Register")
    purch.append(["Date", "Supplier", "Item", "Qty", "Amount", "Paid/Credit"])
    purch.append([date(2026, 6, 25), "ABC Distributors", "Tata Salt", 100, 2200, "PAID"])
    purch.append([date(2026, 6, 26), "XYZ Traders", "New Widget", 20, 400, "CREDIT"])

    exp = wb.create_sheet("Expense Register")
    exp.append(["Date", "Type", "Amount", "Mode", "Notes"])
    exp.append([date(2026, 6, 1), "Electricity", 1500, "UPI", "monthly bill"])
    exp.append([date(2026, 6, 2), "Planned repair", None, "Cash", "not yet paid"])  # no amount -> skipped

    daily = wb.create_sheet("Daily Accounts")
    daily.append(["Date", "Opening", "Morning", "Evening", "UPI In", "Cash Out"])
    daily.append([date(2026, 7, 1), 1000, 2000, 1500, 800, 300])
    daily.append([date(2026, 7, 1), 999, 0, 0, 0, 0])  # duplicate date

    path = tmp_path / "test_workbook.xlsm"
    wb.save(path)
    return path


@pytest.fixture
def workbook_path(tmp_path):
    return _build_workbook(tmp_path)


@pytest.fixture
def loaded_wb(workbook_path):
    import openpyxl as _oxl
    return _oxl.load_workbook(workbook_path, data_only=True)
